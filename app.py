"""News Scanner: search the internet for news, posts and discussions about a topic,
then summarise and chat about what was found.

Stack: Streamlit (UI), Serper.dev (Google search API), LangChain (prompts/chains),
and OpenAI or Replicate as the language model.

Run with:  streamlit run app.py
"""

from __future__ import annotations

import hmac
import io
import os
import re
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date
from typing import Any

import replicate
import requests
import streamlit as st
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from dotenv import load_dotenv
from langchain_core.language_models.llms import LLM
from langchain_core.messages import AIMessage, HumanMessage
from langchain_core.outputs import GenerationChunk
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
from langchain_core.output_parsers import StrOutputParser
from langchain_openai import ChatOpenAI

load_dotenv()

SERPER_URL = "https://google.serper.dev"

# Each source maps to a Serper endpoint plus an optional site filter that is
# appended to the query, so "posts" and "chat" come from the platforms people
# actually post on rather than from the open web.
SOURCES: dict[str, dict[str, str]] = {
    "News articles": {"endpoint": "news", "sites": ""},
    "Social posts": {
        "endpoint": "search",
        "sites": "site:x.com OR site:linkedin.com/posts OR site:facebook.com "
        "OR site:instagram.com OR site:threads.net OR site:bsky.app",
    },
    "Discussions & forums": {
        "endpoint": "search",
        "sites": "site:reddit.com OR site:news.ycombinator.com OR site:quora.com "
        "OR site:stackexchange.com",
    },
    "Chat & messages": {
        "endpoint": "search",
        "sites": "site:t.me OR site:discord.com OR site:mastodon.social OR site:substack.com/notes",
    },
    "Web": {"endpoint": "search", "sites": ""},
}

TIME_RANGES = {
    "Past hour": "qdr:h",
    "Past 24 hours": "qdr:d",
    "Past week": "qdr:w",
    "Past month": "qdr:m",
    "Past year": "qdr:y",
    "Any time": None,
}

# Country name -> (Google country code, local domain suffix)
REGIONS: dict[str, tuple[str | None, str | None]] = {
    "Global": (None, None),
    "Australia": ("au", ".au"),
    "New Zealand": ("nz", ".nz"),
    "United States": ("us", None),
    "United Kingdom": ("uk", ".uk"),
    "Canada": ("ca", ".ca"),
    "Ireland": ("ie", ".ie"),
    "India": ("in", ".in"),
    "Singapore": ("sg", ".sg"),
    "Philippines": ("ph", ".ph"),
    "South Africa": ("za", ".za"),
}


@dataclass(frozen=True)
class Geo:
    """Where to search and how strictly to stick to that region."""

    region: str = "Global"
    must_mention: bool = False  # query must mention the country by name
    local_sites: bool = False  # news/web results only from the country's domains

    @property
    def gl(self) -> str | None:
        return REGIONS[self.region][0]

    @property
    def tld(self) -> str | None:
        return REGIONS[self.region][1] if self.local_sites else None

    @property
    def location(self) -> str | None:
        return None if self.region == "Global" else self.region

    def focus_note(self) -> str:
        if self.region == "Global":
            return ""
        return (
            f" Focus on what is relevant to {self.region}: leave out results that "
            f"aren't, and say so if few results relate to {self.region}."
        )


OPENAI_MODELS = ["gpt-4o-mini", "gpt-4.1-mini", "gpt-4o", "gpt-5-mini"]
REPLICATE_MODELS = [
    "meta/meta-llama-3-8b-instruct",
    "meta/meta-llama-3-70b-instruct",
    "mistralai/mixtral-8x7b-instruct-v0.1",
]


# --------------------------------------------------------------------------- #
# Search (Serper.dev)
# --------------------------------------------------------------------------- #
@dataclass
class Result:
    source: str
    title: str
    link: str
    snippet: str
    date: str = ""
    publisher: str = ""


@st.cache_data(ttl=900, show_spinner=False)
def serper_search(
    query: str,
    endpoint: str,
    num: int,
    tbs: str | None,
    gl: str | None,
    location: str | None,
    _api_key: str,
) -> list[dict[str, Any]]:
    """Call Serper and return the raw result items. Cached for 15 minutes."""
    payload: dict[str, Any] = {"q": query, "num": num}
    for name, value in (("tbs", tbs), ("gl", gl), ("location", location)):
        if value:
            payload[name] = value
    resp = requests.post(
        f"{SERPER_URL}/{endpoint}",
        headers={"X-API-KEY": _api_key, "Content-Type": "application/json"},
        json=payload,
        timeout=20,
    )
    resp.raise_for_status()
    data = resp.json()
    return data.get("news" if endpoint == "news" else "organic", [])


def scan_topic(
    topic: str, sources: list[str], num: int, tbs: str | None, geo: Geo, api_key: str
) -> dict[str, list[Result]]:
    results: dict[str, list[Result]] = {}
    for name in sources:
        cfg = SOURCES[name]
        query = topic
        if geo.must_mention and geo.location:
            query += f' "{geo.location}"'
        # Platform sources keep their own site filter; the local-domain filter
        # only applies to open news/web searches.
        if cfg["sites"]:
            query += f" ({cfg['sites']})"
        elif geo.tld:
            query += f" site:{geo.tld}"
        items = serper_search(query, cfg["endpoint"], num, tbs, geo.gl, geo.location, api_key)
        results[name] = [
            Result(
                source=name,
                title=item.get("title", ""),
                link=item.get("link", ""),
                snippet=item.get("snippet", ""),
                date=item.get("date", ""),
                publisher=item.get("source", ""),
            )
            for item in items
        ]
    return results


def format_context(results: list[Result]) -> str:
    """Number results so the model can cite them as [1], [2], ..."""
    lines = []
    for i, r in enumerate(results, 1):
        meta = " · ".join(x for x in (r.source, r.publisher, r.date) if x)
        lines.append(f"[{i}] {r.title} ({meta})\n{r.snippet}\nURL: {r.link}")
    return "\n\n".join(lines) or "No search results."


# --------------------------------------------------------------------------- #
# Word export
# --------------------------------------------------------------------------- #
INLINE_MD = re.compile(r"(\*\*[^*]+\*\*|\[[^\]]+\]\([^)\s]+\)|\*[^*\s][^*]*\*)")


def add_hyperlink(paragraph, text: str, url: str) -> None:
    """python-docx has no hyperlink API, so build the run XML directly."""
    rel_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    style = OxmlElement("w:rStyle")
    style.set(qn("w:val"), "Hyperlink")
    props.append(style)
    run.append(props)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


def add_markdown_runs(paragraph, text: str) -> None:
    """Render **bold**, *italic* and [links](url) within one line."""
    for part in INLINE_MD.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("[") and "](" in part:
            label, url = part[1:-1].split("](", 1)
            add_hyperlink(paragraph, label, url)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def briefing_to_docx(topic: str, briefing: str, results: list[Result], meta: dict[str, str]) -> bytes:
    doc = Document()
    doc.add_heading(f"Briefing: {topic}", level=0)
    info = doc.add_paragraph()
    info.add_run(" · ".join(f"{k}: {v}" for k, v in meta.items())).italic = True

    for raw in briefing.splitlines():
        line = raw.strip()
        if not line or set(line) <= {"-", "*", "_"}:  # blank or horizontal rule
            continue
        if heading := re.match(r"(#{1,6})\s+(.*)", line):
            doc.add_heading(heading.group(2).strip("* "), level=min(len(heading.group(1)), 3))
        elif bullet := re.match(r"[-*•]\s+(.*)", line):
            level2 = len(raw) - len(raw.lstrip()) >= 2
            add_markdown_runs(doc.add_paragraph(style="List Bullet 2" if level2 else "List Bullet"), bullet.group(1))
        elif numbered := re.match(r"\d+[.)]\s+(.*)", line):
            add_markdown_runs(doc.add_paragraph(style="List Number"), numbered.group(1))
        else:
            add_markdown_runs(doc.add_paragraph(), line)

    doc.add_heading("Sources", level=1)
    for i, r in enumerate(results, 1):
        p = doc.add_paragraph()
        p.add_run(f"[{i}] ")
        add_hyperlink(p, r.title or r.link, r.link)
        details = " · ".join(x for x in (r.source, r.publisher, r.date) if x)
        if details:
            p.add_run(f" — {details}").italic = True

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
# Language models
# --------------------------------------------------------------------------- #
class ReplicateLLM(LLM):
    """Minimal LangChain wrapper around a Replicate-hosted text model."""

    model: str
    api_token: str
    temperature: float = 0.3
    max_tokens: int = 1024

    @property
    def _llm_type(self) -> str:
        return "replicate"

    def _input(self, prompt: str) -> dict[str, Any]:
        return {
            "prompt": prompt,
            "temperature": self.temperature,
            "max_tokens": self.max_tokens,
            "max_new_tokens": self.max_tokens,
        }

    def _stream(self, prompt: str, stop=None, run_manager=None, **kwargs) -> Iterator[GenerationChunk]:
        client = replicate.Client(api_token=self.api_token)
        for event in client.stream(self.model, input=self._input(prompt)):
            text = str(event)
            if text:
                if run_manager:
                    run_manager.on_llm_new_token(text)
                yield GenerationChunk(text=text)

    def _call(self, prompt: str, stop=None, run_manager=None, **kwargs) -> str:
        return "".join(chunk.text for chunk in self._stream(prompt, stop, run_manager))


def build_llm(provider: str, model: str, keys: dict[str, str]):
    if provider == "OpenAI":
        return ChatOpenAI(model=model, api_key=keys["openai"])
    return ReplicateLLM(model=model, api_token=keys["replicate"])


BRIEFING_PROMPT = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "You are a news analyst. Using ONLY the numbered search results, write a "
            "concise briefing on the topic. Include: a 2-3 sentence overview, key "
            "developments as bullet points, what people are saying on social media "
            "and forums (overall sentiment), and notable open questions. Cite sources "
            "inline like [3]. If results are thin or conflicting, say so.{focus}",
        ),
        ("human", "Topic: {topic}\n\nSearch results:\n{context}"),
    ]
)

CHAT_PROMPT = ChatPromptTemplate.from_messages(
    [
        (
            "system",
            "You answer questions about the topic '{topic}' using the numbered search "
            "results below, which were just retrieved from the internet. Cite sources "
            "like [2]. If the results don't answer the question, say so plainly.{focus}\n\n"
            "Search results:\n{context}",
        ),
        MessagesPlaceholder("history"),
        ("human", "{question}"),
    ]
)


# --------------------------------------------------------------------------- #
# UI
# --------------------------------------------------------------------------- #
def get_secret(name: str) -> str:
    """Read a value from the environment / .env, then Streamlit secrets."""
    value = os.getenv(name, "")
    try:
        value = value or st.secrets.get(name, "")
    except Exception:  # no secrets.toml
        pass
    return value


def get_key(label: str, env_var: str) -> str:
    # A configured key never goes into a widget: widget values are sent to the
    # browser, so a public app would leak it to every visitor.
    if configured := get_secret(env_var):
        st.caption(f"✅ {label} configured")
        return configured
    return st.text_input(label, type="password", key=env_var)


def check_password() -> bool:
    """Gate the app behind APP_PASSWORD if one is configured."""
    password = get_secret("APP_PASSWORD")
    if not password or st.session_state.get("authenticated"):
        return True
    entered = st.text_input("Password", type="password")
    if entered and hmac.compare_digest(entered, password):
        st.session_state.authenticated = True
        st.rerun()
    elif entered:
        st.error("Incorrect password.")
    return False


def sidebar() -> dict[str, Any]:
    with st.sidebar:
        st.header("Settings")
        provider = st.radio("Language model", ["OpenAI", "Replicate"], horizontal=True)
        models = OPENAI_MODELS if provider == "OpenAI" else REPLICATE_MODELS
        model = st.selectbox("Model", models, accept_new_options=True)

        st.subheader("API keys")
        keys = {"serper": get_key("Serper.dev API key", "SERPER_API_KEY")}
        if provider == "OpenAI":
            keys["openai"] = get_key("OpenAI API key", "OPENAI_API_KEY")
        else:
            keys["replicate"] = get_key("Replicate API token", "REPLICATE_API_TOKEN")

        st.subheader("Search")
        time_range = st.selectbox("Time range", list(TIME_RANGES), index=2)
        num = st.slider("Results per source", 5, 20, 10)

        st.subheader("Region")
        region = st.selectbox("Search from", list(REGIONS), index=1)
        must_mention = local_sites = False
        if region != "Global":
            must_mention = st.checkbox(
                f"Only results mentioning {region}",
                help=f'Adds "{region}" to every search so results must be about it.',
            )
            if REGIONS[region][1]:
                local_sites = st.checkbox(
                    f"Only {REGIONS[region][1]} websites (news & web)",
                    help="Social, forum and chat sources keep their platform filters.",
                )
            st.caption(f"Results are ranked as seen from {region}, and the AI focuses on what's relevant there.")

    return {
        "provider": provider,
        "model": model,
        "keys": keys,
        "tbs": TIME_RANGES[time_range],
        "time_range": time_range,
        "num": num,
        "geo": Geo(region, must_mention, local_sites),
    }


def render_results(results: dict[str, list[Result]]) -> None:
    tabs = st.tabs([f"{name} ({len(items)})" for name, items in results.items()], key="source_tabs")
    for tab, items in zip(tabs, results.values()):
        with tab:
            if not items:
                st.caption("No results for this source.")
            for r in items:
                meta = " · ".join(x for x in (r.publisher, r.date) if x)
                st.markdown(f"**[{r.title}]({r.link})**")
                if meta:
                    st.caption(meta)
                st.write(r.snippet)
                st.divider()


def main() -> None:
    st.set_page_config(page_title="News Scanner", page_icon="📰", layout="wide")
    st.title("📰 News Scanner")
    st.caption("Search news, social posts, forums and chat channels about any topic, then ask questions about it.")

    if not check_password():
        return

    cfg = sidebar()
    state = st.session_state
    state.setdefault("results", {})
    state.setdefault("topic", "")
    state.setdefault("briefing", "")
    state.setdefault("messages", [])
    state.setdefault("scan_meta", {})

    with st.form("scan"):
        col1, col2 = st.columns([3, 2])
        topic = col1.text_input("Topic", placeholder="e.g. housing affordability")
        sources = col2.multiselect("Sources", list(SOURCES), default=list(SOURCES)[:4])
        submitted = st.form_submit_button("Scan the internet", type="primary")

    llm_key = cfg["keys"].get("openai") or cfg["keys"].get("replicate")
    if submitted:
        if not topic.strip() or not sources:
            st.warning("Enter a topic and choose at least one source.")
        elif not cfg["keys"]["serper"]:
            st.error("Add your Serper.dev API key in the sidebar.")
        else:
            try:
                with st.spinner(f"Searching for “{topic}”…"):
                    state.results = scan_topic(
                        topic, sources, cfg["num"], cfg["tbs"], cfg["geo"], cfg["keys"]["serper"]
                    )
                state.topic, state.briefing, state.messages = topic, "", []
                # Remember the scan's own settings so later sidebar changes don't
                # alter the briefing or export for results already fetched.
                state.geo = cfg["geo"]
                state.scan_meta = {
                    "Date": date.today().strftime("%d %B %Y"),
                    "Region": cfg["geo"].region,
                    "Time range": cfg["time_range"],
                    "Sources": ", ".join(sources),
                }
            except requests.HTTPError as e:
                st.error(f"Search failed: {e}")

    if not state.results:
        st.info("Enter a topic and click **Scan the internet** to get started.")
        return

    all_results = [r for items in state.results.values() for r in items]
    briefing_tab, results_tab, chat_tab = st.tabs(["🧠 Briefing", "🔎 Results", "💬 Chat"], key="main_tabs")

    with results_tab:
        render_results(state.results)

    with briefing_tab:
        if not llm_key:
            st.warning(f"Add your {cfg['provider']} key in the sidebar to generate a briefing.")
        elif state.briefing:
            st.markdown(state.briefing)
            slug = re.sub(r"[^a-z0-9]+", "-", state.topic.lower()).strip("-") or "topic"
            st.download_button(
                "⬇️ Download as Word (.docx)",
                data=briefing_to_docx(state.topic, state.briefing, all_results, state.scan_meta),
                file_name=f"briefing-{slug}-{date.today():%Y-%m-%d}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary",
            )
        elif st.button("Generate briefing", type="primary"):
            chain = BRIEFING_PROMPT | build_llm(cfg["provider"], cfg["model"], cfg["keys"]) | StrOutputParser()
            try:
                state.briefing = st.write_stream(
                    chain.stream(
                        {"topic": state.topic, "context": format_context(all_results), "focus": state.geo.focus_note()}
                    )
                )
                st.rerun()  # redraw with the download button
            except Exception as e:
                st.error(f"Model call failed: {e}")
        with st.expander("Sources used"):
            for i, r in enumerate(all_results, 1):
                st.markdown(f"[{i}] [{r.title}]({r.link}) — {r.source}")

    with chat_tab:
        st.caption(f"Ask about **{state.topic}**. Each question runs a fresh search too.")
        for m in state.messages:
            with st.chat_message("user" if isinstance(m, HumanMessage) else "assistant"):
                st.markdown(m.content)

    question = st.chat_input(f"Ask about {state.topic}…", disabled=not llm_key)
    if question:
        with chat_tab:
            with st.chat_message("user"):
                st.markdown(question)
            with st.chat_message("assistant"):
                try:
                    with st.spinner("Searching…"):
                        fresh = scan_topic(
                            f"{state.topic} {question}",
                            ["News articles", "Web"],
                            5,
                            cfg["tbs"],
                            state.geo,
                            cfg["keys"]["serper"],
                        )
                    context = format_context([r for items in fresh.values() for r in items] + all_results)
                    chain = CHAT_PROMPT | build_llm(cfg["provider"], cfg["model"], cfg["keys"]) | StrOutputParser()
                    answer = st.write_stream(
                        chain.stream(
                            {
                                "topic": state.topic,
                                "focus": state.geo.focus_note(),
                                "context": context,
                                "history": state.messages[-10:],
                                "question": question,
                            }
                        )
                    )
                    state.messages += [HumanMessage(question), AIMessage(answer)]
                except Exception as e:
                    st.error(f"Something went wrong: {e}")


if __name__ == "__main__":
    main()
